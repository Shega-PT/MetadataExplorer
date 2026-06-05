#!/usr/bin/env python3
"""
Metadata Explorer v2 — Universal Metadata Extractor
Copyright (C) 2026  ShegaPT

This program is free software: you can redistribute it and/or modify
it under the terms of the GNU General Public License as published by
the Free Software Foundation, either version 3 of the License, or
(at your option) any later version.

This program is distributed in the hope that it will be useful,
but WITHOUT ANY WARRANTY; without even the implied warranty of
MERCHANTABILITY or FITNESS FOR A PARTICULAR PURPOSE.  See the
GNU General Public License for more details.

You should have received a copy of the GNU General Public License
along with this program.  If not, see <https://www.gnu.org/licenses/>.
"""

import argparse
import csv
import json
import logging
import os
import subprocess
import sys
import warnings
from concurrent.futures import ThreadPoolExecutor, as_completed
from datetime import datetime
from pathlib import Path
from typing import Dict, Any, Optional, List, Tuple

import argparse
import csv
import json
import logging
import os
import subprocess
import sys
import warnings
from concurrent.futures import ThreadPoolExecutor, as_completed
from datetime import datetime
from pathlib import Path
from typing import Dict, Any, Optional, List, Tuple

# Suppress noisy library warnings
warnings.filterwarnings("ignore", category=UserWarning)
logging.getLogger("exifread").setLevel(logging.ERROR)
logging.getLogger("pypdf").setLevel(logging.ERROR)
logging.getLogger("PyPDF2").setLevel(logging.ERROR)
logging.getLogger("PIL").setLevel(logging.WARNING)

# ============================================================================
# Optional Imports with Graceful Degradation
# ============================================================================

try:
    import exifread
except ImportError:
    exifread = None

try:
    import mutagen
except ImportError:
    mutagen = None

try:
    from PIL import Image

    try:
        from pillow_heif import register_heif_opener

        register_heif_opener()
    except ImportError:
        pass
except ImportError:
    Image = None

try:
    from hachoir.metadata import extractMetadata
    from hachoir.parser import createParser
except ImportError:
    extractMetadata = None
    createParser = None

try:
    from pymediainfo import MediaInfo
except ImportError:
    MediaInfo = None

try:
    import magic
except ImportError:
    magic = None

try:
    import pypdf
except ImportError:
    try:
        import PyPDF2 as pypdf
    except ImportError:
        pypdf = None

try:
    import docx
except ImportError:
    docx = None

try:
    import openpyxl
except ImportError:
    openpyxl = None

try:
    from tqdm import tqdm
except ImportError:
    tqdm = None


# ============================================================================
# LOGGING
# ============================================================================


def setup_logging(verbose: bool = False) -> logging.Logger:
    logger = logging.getLogger("universal_metadata")
    logger.setLevel(logging.DEBUG if verbose else logging.INFO)
    if not logger.handlers:
        handler = logging.StreamHandler()
        handler.setFormatter(
            logging.Formatter("%(asctime)s - %(levelname)s - %(message)s")
        )
        logger.addHandler(handler)
    return logger


# ============================================================================
# UNIVERSAL METADATA EXTRACTOR
# ============================================================================


class UniversalMetadataExtractor:
    IMAGE_EXTENSIONS = {
        ".jpg",
        ".jpeg",
        ".png",
        ".tiff",
        ".tif",
        ".webp",
        ".heic",
        ".heif",
        ".bmp",
        ".gif",
    }
    AUDIO_EXTENSIONS = {
        ".mp3",
        ".flac",
        ".m4a",
        ".ogg",
        ".wav",
        ".aac",
        ".wma",
        ".opus",
    }
    VIDEO_EXTENSIONS = {".mp4", ".mov", ".avi", ".mkv", ".wmv", ".flv", ".webm", ".m4v"}
    DOCUMENT_EXTENSIONS = {".pdf", ".docx", ".xlsx"}
    RAW_EXTENSIONS = {
        ".cr2",
        ".cr3",
        ".nef",
        ".nrw",
        ".arw",
        ".dng",
        ".orf",
        ".rw2",
        ".raf",
    }
    ALL_EXTENSIONS = (
        IMAGE_EXTENSIONS
        | AUDIO_EXTENSIONS
        | VIDEO_EXTENSIONS
        | DOCUMENT_EXTENSIONS
        | RAW_EXTENSIONS
    )

    TYPE_FILTERS = {
        "images": IMAGE_EXTENSIONS | RAW_EXTENSIONS,
        "audio": AUDIO_EXTENSIONS,
        "video": VIDEO_EXTENSIONS,
        "documents": DOCUMENT_EXTENSIONS,
        "all": None,
    }

    @staticmethod
    def _format_size(size_bytes: int) -> str:
        for unit in ("B", "KB", "MB", "GB", "TB"):
            if size_bytes < 1024:
                return f"{size_bytes:.2f} {unit}" if unit != "B" else f"{size_bytes} B"
            size_bytes /= 1024
        return f"{size_bytes:.2f} PB"

    @staticmethod
    def _gps_to_decimal(gps_list, ref: str) -> Optional[float]:
        try:
            degrees = float(gps_list[0])
            minutes = float(gps_list[1])
            seconds = float(gps_list[2])
            decimal = degrees + minutes / 60.0 + seconds / 3600.0
            return -decimal if ref in ("S", "W") else decimal
        except (TypeError, IndexError, ValueError):
            return None

    @staticmethod
    def _seconds_to_hms(seconds: float) -> str:
        h, r = divmod(int(seconds), 3600)
        m, s = divmod(r, 60)
        parts = []
        if h:
            parts.append(f"{h}h")
        if m:
            parts.append(f"{m}m")
        parts.append(f"{s}s")
        return " ".join(parts)

    # ------------------------------------------------------------------
    # Image metadata
    # ------------------------------------------------------------------
    @classmethod
    def get_image_metadata(cls, file_path: Path) -> Dict[str, Any]:
        meta = {}

        if Image is not None:
            try:
                with Image.open(file_path) as img:
                    meta["IMG_Format"] = img.format or "Unknown"
                    meta["IMG_Width"] = str(img.width)
                    meta["IMG_Height"] = str(img.height)
                    meta["IMG_Dimensions"] = f"{img.width}x{img.height}"
                    meta["IMG_Mode"] = img.mode
                    dpi = img.info.get("dpi")
                    if dpi and isinstance(dpi, tuple) and len(dpi) == 2:
                        meta["IMG_DPI"] = f"{dpi[0]:.0f}x{dpi[1]:.0f}"
            except Exception:
                pass

        if exifread is not None:
            try:
                with open(file_path, "rb") as f:
                    tags = exifread.process_file(f, details=True)
                for tag, value in tags.items():
                    if tag in ("JPEGThumbnail", "TIFFThumbnail", "Filename"):
                        continue
                    tag_str = str(value)
                    if len(tag_str) > 1000:
                        continue
                    key = f"IMG_{tag.replace(' ', '_')}"

                    if tag == "GPS GPSLatitude":
                        ref = str(tags.get("GPS GPSLatitudeRef", "N"))
                        dec = cls._gps_to_decimal(value.values, ref)
                        if dec is not None:
                            meta["IMG_GPS_LatitudeDecimal"] = f"{dec:.6f}"
                        meta[key] = tag_str
                    elif tag == "GPS GPSLongitude":
                        ref = str(tags.get("GPS GPSLongitudeRef", "E"))
                        dec = cls._gps_to_decimal(value.values, ref)
                        if dec is not None:
                            meta["IMG_GPS_LongitudeDecimal"] = f"{dec:.6f}"
                        meta[key] = tag_str
                    else:
                        meta[key] = tag_str
            except Exception:
                pass

        return meta

    # ------------------------------------------------------------------
    # Audio metadata
    # ------------------------------------------------------------------
    @classmethod
    def get_audio_metadata(cls, file_path: Path) -> Dict[str, Any]:
        meta = {}
        if mutagen is None:
            return meta

        try:
            audio = mutagen.File(file_path)
            if audio is None:
                return meta

            if hasattr(audio, "info"):
                info = audio.info
                if hasattr(info, "length") and info.length:
                    meta["AUDIO_Duration"] = cls._seconds_to_hms(info.length)
                    meta["AUDIO_DurationSeconds"] = f"{info.length:.2f}"
                if hasattr(info, "bitrate") and info.bitrate:
                    br = info.bitrate
                    meta["AUDIO_Bitrate"] = (
                        f"{br // 1000} kbps" if br >= 1000 else f"{br} bps"
                    )
                if hasattr(info, "sample_rate") and info.sample_rate:
                    meta["AUDIO_SampleRate"] = f"{info.sample_rate} Hz"
                if hasattr(info, "channels") and info.channels:
                    meta["AUDIO_Channels"] = str(info.channels)

            if audio.tags:
                for tag_name in audio.tags:
                    try:
                        values = audio.tags[tag_name]
                        raw = values[0] if isinstance(values, list) else values
                        value = str(raw)
                        if value:
                            meta[f"AUDIO_{tag_name}"] = value
                    except Exception:
                        continue
        except Exception:
            pass

        return meta

    # ------------------------------------------------------------------
    # Video metadata (pymediainfo -> hachoir -> ffprobe)
    # ------------------------------------------------------------------
    @classmethod
    def get_video_metadata(cls, file_path: Path) -> Dict[str, Any]:
        meta = {}

        if MediaInfo is not None:
            try:
                mi = MediaInfo.parse(str(file_path))
                for track in mi.tracks:
                    prefix = f"VIDEO_{track.track_type.upper()}"
                    for k, v in track.to_data().items():
                        if v is not None and k != "track_type":
                            clean_k = k.replace(" ", "_").replace("/", "_")
                            meta[f"{prefix}_{clean_k}"] = str(v)
                return meta
            except Exception:
                pass

        if extractMetadata is not None and createParser is not None:
            try:
                parser = createParser(str(file_path))
                if parser:
                    with parser:
                        extracted = extractMetadata(parser)
                        if extracted:
                            for line in extracted.exportPlaintext():
                                if ":" in line:
                                    k, v = line.split(":", 1)
                                    meta[f"VIDEO_{k.strip().replace(' ', '_')}"] = (
                                        v.strip()
                                    )
                    return meta
            except Exception:
                pass

        try:
            result = subprocess.run(
                [
                    "ffprobe",
                    "-v",
                    "quiet",
                    "-print_format",
                    "json",
                    "-show_format",
                    "-show_streams",
                    str(file_path),
                ],
                capture_output=True,
                text=True,
                timeout=30,
            )
            if result.returncode == 0:
                data = json.loads(result.stdout)
                if "format" in data:
                    for k, v in data["format"].items():
                        if v is not None:
                            meta[f"VIDEO_FORMAT_{k}"] = str(v)
                if "streams" in data:
                    for i, stream in enumerate(data["streams"]):
                        for k, v in stream.items():
                            if v is not None and k != "disposition":
                                meta[f"VIDEO_Stream{i}_{k}"] = str(v)
        except Exception:
            pass

        return meta

    # ------------------------------------------------------------------
    # Document metadata
    # ------------------------------------------------------------------
    @staticmethod
    def get_document_metadata(file_path: Path) -> Dict[str, Any]:
        meta = {}
        ext = file_path.suffix.lower()

        try:
            if ext == ".pdf" and pypdf is not None:
                with open(file_path, "rb") as f:
                    reader = pypdf.PdfReader(f)
                meta["DOC_Pages"] = str(len(reader.pages))
                info = reader.metadata
                if info:
                    for k, v in info.items():
                        if v:
                            key = k.replace("/", "_").strip()
                            meta[f"DOC_{key}"] = str(v)

            elif ext == ".docx" and docx is not None:
                document = docx.Document(str(file_path))
                props = document.core_properties
                if props.author:
                    meta["DOC_Author"] = props.author
                if props.title:
                    meta["DOC_Title"] = props.title
                if props.created:
                    meta["DOC_Created"] = str(props.created)
                if props.modified:
                    meta["DOC_Modified"] = str(props.modified)
                if props.subject:
                    meta["DOC_Subject"] = props.subject

            elif ext == ".xlsx" and openpyxl is not None:
                wb = openpyxl.load_workbook(file_path, read_only=True, data_only=True)
                props = wb.properties
                meta["DOC_Sheets"] = str(len(wb.sheetnames))
                meta["DOC_SheetNames"] = ", ".join(wb.sheetnames[:20])
                if len(wb.sheetnames) > 20:
                    meta["DOC_SheetNames"] += f" ... (+{len(wb.sheetnames) - 20} more)"
                if props.creator:
                    meta["DOC_Author"] = props.creator
                if props.title:
                    meta["DOC_Title"] = props.title
                if props.created:
                    meta["DOC_Created"] = str(props.created)
                if props.modified:
                    meta["DOC_Modified"] = str(props.modified)
                wb.close()
        except Exception:
            pass

        return meta

    # ------------------------------------------------------------------
    # Basic file metadata
    # ------------------------------------------------------------------
    @staticmethod
    def get_basic_metadata(file_path: Path) -> Dict[str, Any]:
        meta = {}
        try:
            st = file_path.stat()
            meta["FILE_Size"] = UniversalMetadataExtractor._format_size(st.st_size)
            meta["FILE_SizeBytes"] = str(st.st_size)
            meta["FILE_Created"] = datetime.fromtimestamp(st.st_ctime).isoformat()
            meta["FILE_Modified"] = datetime.fromtimestamp(st.st_mtime).isoformat()
            meta["FILE_Extension"] = file_path.suffix.lower()
        except Exception:
            pass
        return meta

    # ------------------------------------------------------------------
    # Detection & dispatch
    # ------------------------------------------------------------------
    @staticmethod
    def detect_mime_type(file_path: Path) -> Optional[str]:
        if magic is not None:
            try:
                return magic.from_file(str(file_path), mime=True)
            except Exception:
                pass
        return None

    @classmethod
    def classify_extension(cls, file_path: Path) -> str:
        ext = file_path.suffix.lower()
        for cat, exts in (
            ("image", cls.IMAGE_EXTENSIONS),
            ("audio", cls.AUDIO_EXTENSIONS),
            ("video", cls.VIDEO_EXTENSIONS),
            ("document", cls.DOCUMENT_EXTENSIONS),
            ("raw", cls.RAW_EXTENSIONS),
        ):
            if ext in exts:
                return cat
        return "unknown"

    @classmethod
    def get_all_metadata(cls, file_path: Path) -> Dict[str, Any]:
        meta = {}
        ext = file_path.suffix.lower()

        mime = cls.detect_mime_type(file_path)
        if mime:
            meta["FILE_MimeType"] = mime

        if ext in cls.IMAGE_EXTENSIONS or ext in cls.RAW_EXTENSIONS:
            meta.update(cls.get_image_metadata(file_path))
        elif ext in cls.AUDIO_EXTENSIONS:
            meta.update(cls.get_audio_metadata(file_path))
        elif ext in cls.VIDEO_EXTENSIONS:
            meta.update(cls.get_video_metadata(file_path))
        elif ext in cls.DOCUMENT_EXTENSIONS:
            meta.update(cls.get_document_metadata(file_path))

        meta.update(cls.get_basic_metadata(file_path))
        return meta


# ============================================================================
# OUTPUT WRITERS
# ============================================================================


def write_log_output(
    results: List[Tuple[Path, str, Dict[str, Any]]], output_path: str
) -> None:
    results.sort(key=lambda r: (r[0].parent, r[1]))
    with open(output_path, "w", encoding="utf-8") as f:
        current_parent = None
        for rel_path, filename, metadata in results:
            parent = rel_path.parent
            if parent != current_parent:
                f.write(f"\n{'=' * 80}\n")
                label = (
                    "ROOT DIRECTORY" if str(parent) == "." else f"DIRECTORY: {parent}"
                )
                f.write(f"{label}\n")
                f.write(f"{'=' * 80}\n")
                current_parent = parent

            f.write(f"\nFILE: {filename}\n")
            f.write(f"PATH: {rel_path}\n")
            if metadata:
                f.write("METADATA:\n")
                for key in sorted(metadata.keys()):
                    value = str(metadata[key])
                    if len(value) > 500:
                        value = value[:497] + "..."
                    f.write(f"  \u2022 {key}: {value}\n")
            else:
                f.write("  \u2022 No extractable metadata found\n")


def write_json_output(
    results: List[Tuple[Path, str, Dict[str, Any]]], output_path: str
) -> None:
    data = [
        {"file": filename, "path": str(rel_path), "metadata": metadata}
        for rel_path, filename, metadata in results
    ]
    with open(output_path, "w", encoding="utf-8") as f:
        json.dump(data, f, indent=2, ensure_ascii=False)


def write_csv_output(
    results: List[Tuple[Path, str, Dict[str, Any]]], output_path: str
) -> None:
    all_keys = sorted({k for _, _, meta in results for k in meta})
    with open(output_path, "w", encoding="utf-8", newline="") as f:
        writer = csv.writer(f)
        writer.writerow(["File", "Path"] + all_keys)
        for rel_path, filename, metadata in results:
            writer.writerow(
                [filename, str(rel_path)] + [metadata.get(k, "") for k in all_keys]
            )


# ============================================================================
# FILE MANAGER
# ============================================================================


class FileManager:
    def __init__(
        self,
        base_folder: str,
        file_filter: str = "all",
        dry_run: bool = False,
        num_threads: int = 1,
        show_progress: bool = False,
    ):
        self.base_folder = Path(base_folder).resolve()
        self.dry_run = dry_run
        self.num_threads = (os.cpu_count() or 1) if num_threads <= 0 else num_threads
        self.show_progress = show_progress
        self.ignored_dirs = {
            ".git",
            "__pycache__",
            ".venv",
            "node_modules",
            ".mypy_cache",
            ".pytest_cache",
        }
        self.ignored_files = {".DS_Store", "Thumbs.db", "desktop.ini"}
        self.valid_extensions = UniversalMetadataExtractor.TYPE_FILTERS.get(file_filter)

    def _should_skip(self, path: Path) -> bool:
        if any(part.startswith(".") for part in path.parts):
            return True
        if path.is_dir() and path.name in self.ignored_dirs:
            return True
        if path.is_file() and path.name in self.ignored_files:
            return True
        return False

    def _collect_files(self) -> List[Path]:
        files = []
        for root, dirs, filenames in os.walk(self.base_folder, topdown=True):
            current = Path(root)
            dirs[:] = [d for d in dirs if not self._should_skip(current / d)]
            if self._should_skip(current):
                continue
            for name in sorted(filenames):
                fp = current / name
                if self._should_skip(fp):
                    continue
                if name == Path(__file__).name or name.endswith(
                    (".log", ".json", ".csv")
                ):
                    continue
                if (
                    self.valid_extensions is not None
                    and fp.suffix.lower() not in self.valid_extensions
                ):
                    continue
                files.append(fp)
        return files

    def _process_file(self, file_path: Path) -> Tuple[Path, str, Dict[str, Any]]:
        rel = file_path.relative_to(self.base_folder)
        if self.dry_run:
            return rel, file_path.name, {}
        return (
            rel,
            file_path.name,
            UniversalMetadataExtractor.get_all_metadata(file_path),
        )

    def run(self, log: logging.Logger) -> List[Tuple[Path, str, Dict[str, Any]]]:
        log.info(f"Starting exploration of: {self.base_folder}")
        if self.dry_run:
            log.info("DRY RUN MODE \u2013 no metadata will be extracted")

        log.info("Scanning directory for supported files...")
        files = self._collect_files()
        total = len(files)
        log.info(f"Found {total} supported file(s)")

        if total == 0:
            return []

        if self.dry_run:
            log.info("Files that would be processed:")
            for f in files:
                log.info(f"  \u2022 {f.relative_to(self.base_folder)}")
            return []

        results: List[Tuple[Path, str, Dict[str, Any]]] = []
        progress = None

        if self.show_progress and tqdm is not None:
            progress = tqdm(total=total, unit="file", desc="Extracting")

        if self.num_threads > 1:
            with ThreadPoolExecutor(max_workers=self.num_threads) as pool:
                futs = {pool.submit(self._process_file, f): f for f in files}
                for future in as_completed(futs):
                    f = futs[future]
                    try:
                        results.append(future.result())
                    except Exception as e:
                        rel = f.relative_to(self.base_folder)
                        log.warning(f"Error processing {rel}: {e}")
                        results.append((rel, f.name, {}))
                    if progress:
                        progress.update(1)
        else:
            for fp in files:
                try:
                    results.append(self._process_file(fp))
                except Exception as e:
                    rel = fp.relative_to(self.base_folder)
                    log.warning(f"Error processing {rel}: {e}")
                    results.append((rel, fp.name, {}))
                if progress:
                    progress.update(1)

        if progress:
            progress.close()

        supported = sum(1 for _, _, m in results if m)
        log.info(f"\n{'=' * 60}")
        log.info("SCAN COMPLETE")
        log.info(f"{'=' * 60}")
        log.info(f"Files processed: {len(results)}")
        log.info(f"Files with metadata: {supported}")

        return results


# ============================================================================
# CLI
# ============================================================================


DEFAULT_SUBFOLDER = "files"


def parse_args(argv: Optional[List[str]] = None) -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        prog="universal",
        description="Universal Metadata Extractor \u2013 extract metadata from images, audio, video, and documents",
        formatter_class=argparse.RawDescriptionHelpFormatter,
        epilog="""Examples:
  %(prog)s                                          Scan 'files/' subfolder or prompt
  %(prog)s ~/Pictures                               Scan specific directory
  %(prog)s files --format json                      JSON output
  %(prog)s files --format csv -o report.csv          CSV output
  %(prog)s files --type images                      Images only
  %(prog)s files --progress --threads 4             Progress + 4 threads
  %(prog)s files --dry-run                          List files only
""",
    )
    parser.add_argument(
        "directory",
        nargs="?",
        default=None,
        help=f"Directory to scan (default: looks for '{DEFAULT_SUBFOLDER}/' subfolder, or prompts interactively)",
    )
    parser.add_argument(
        "-o", "--output", help="Output file path (default: metadata_report.{format})"
    )
    parser.add_argument(
        "-f",
        "--format",
        choices=["log", "json", "csv"],
        default="log",
        help="Output format (default: log)",
    )
    parser.add_argument(
        "-t",
        "--type",
        choices=["all", "images", "audio", "video", "documents"],
        default="all",
        help="Filter by file type (default: all)",
    )
    parser.add_argument(
        "--dry-run", action="store_true", help="List files without extracting metadata"
    )
    parser.add_argument(
        "-p", "--progress", action="store_true", help="Show progress bar"
    )
    parser.add_argument(
        "-n",
        "--threads",
        type=int,
        default=1,
        help="Number of worker threads (0=auto, default: 1)",
    )
    parser.add_argument(
        "-v", "--verbose", action="store_true", help="Verbose debug output"
    )
    return parser.parse_args(argv)


# ============================================================================
# MAIN
# ============================================================================


def resolve_target_directory(args: argparse.Namespace, log: logging.Logger) -> Path:
    if args.directory is not None:
        target = Path(args.directory)
        if not target.exists():
            log.error(f"Directory does not exist: {args.directory}")
            sys.exit(1)
        if not target.is_dir():
            log.error(f"Not a directory: {args.directory}")
            sys.exit(1)
        return target

    subfolder = Path(DEFAULT_SUBFOLDER)
    if subfolder.exists() and subfolder.is_dir() and any(subfolder.iterdir()):
        log.info(f"Found '{DEFAULT_SUBFOLDER}/' subfolder with content")
        return subfolder

    while True:
        raw = input(
            f"No '{DEFAULT_SUBFOLDER}/' subfolder found. Enter directory path to scan (or 'q' to quit): "
        ).strip()
        if raw.lower() in ("q", "quit", ""):
            log.info("Aborted by user")
            sys.exit(0)
        target = Path(raw).resolve()
        if not target.exists():
            print(f"  Path does not exist: {raw}")
            continue
        if not target.is_dir():
            print(f"  Not a directory: {raw}")
            continue
        return target


def main(argv: Optional[List[str]] = None) -> None:
    args = parse_args(argv)
    log = setup_logging(verbose=args.verbose)

    target = resolve_target_directory(args, log)

    output = args.output or f"metadata_report.{args.format}"

    manager = FileManager(
        base_folder=str(target),
        file_filter=args.type,
        dry_run=args.dry_run,
        num_threads=args.threads,
        show_progress=args.progress,
    )

    try:
        results = manager.run(log)

        if results and not args.dry_run:
            log.info(f"Writing {args.format.upper()} report to {output} ...")
            if args.format == "log":
                write_log_output(results, output)
            elif args.format == "json":
                write_json_output(results, output)
            else:
                write_csv_output(results, output)
            log.info("Done!")
    except KeyboardInterrupt:
        log.warning("Interrupted by user")
        sys.exit(130)
    except Exception as e:
        log.exception(f"Fatal error: {e}")
        sys.exit(1)


if __name__ == "__main__":
    main()
